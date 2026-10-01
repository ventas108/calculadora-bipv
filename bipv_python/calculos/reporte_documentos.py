# -*- coding: utf-8 -*-
"""📄 Reporte en Word (.docx editable) y PDF a partir del mismo HTML.

Spec ``07-informes/reporte-word-pdf`` (1-oct-2026). Abrir el HTML del reporte
en Word perdía las gráficas (Word no lee SVG) y el PDF dependía de imprimir
desde el navegador. Aquí el HTML se recorre una sola vez y se convierte en
bloques (títulos, párrafos, notas, listas, tablas e imágenes); con esos
bloques se arman un .docx editable (python-docx) y un PDF (fpdf2). Las
gráficas SVG pasan a PNG con ``calculos.svg_a_png``. Nada se recalcula: el
contenido es exactamente el del reporte. Módulo puro: sin Streamlit.
"""
from __future__ import annotations

import io
import os
import re
from typing import Any

from lxml import etree, html as lhtml

from calculos.svg_a_png import svg_a_png

_BLOQUES = {"div", "p", "section", "table", "ul", "ol", "h1", "h2", "h3", "h4", "svg", "pre"}
_OMITIR = {"head", "style", "script", "title", "meta", "link", "noscript"}
_ESTADOS = {"🟢": (46, 125, 50), "🟠": (230, 126, 34), "🔴": (198, 40, 40), "🟡": (183, 149, 11),
            "✅": (46, 125, 50), "❌": (198, 40, 40)}


def _limpio(texto: str | None) -> str:
    return re.sub(r"\s+", " ", texto or "").strip()


def _estilo(el, prop: str) -> str | None:
    m = re.search(rf"(?:^|;)\s*{prop}\s*:\s*([^;]+)", el.get("style") or "")
    return m.group(1).strip() if m else None


def _hex(color: str | None) -> str | None:
    if not color:
        return None
    c = color.strip().lower()
    if re.fullmatch(r"#[0-9a-f]{6}", c):
        return c[1:].upper()
    if re.fullmatch(r"#[0-9a-f]{3}", c):
        return "".join(ch * 2 for ch in c[1:]).upper()
    if c == "white":
        return "FFFFFF"
    return None


def _es_oscuro(hexa: str | None) -> bool:
    if not hexa:
        return False
    r, g, b = (int(hexa[i:i + 2], 16) for i in (0, 2, 4))
    return 0.299 * r + 0.587 * g + 0.114 * b < 140


def _clases(el) -> set[str]:
    return set((el.get("class") or "").split())


def bloques_desde_html(html: str) -> list[dict]:
    """El reporte HTML como lista de bloques en orden de lectura."""
    doc = lhtml.fromstring(html)
    cuerpo = doc.find(".//body")
    raiz = cuerpo if cuerpo is not None else doc
    out: list[dict] = []

    def parrafo(texto, fondo=None, negrita=False):
        t = _limpio(texto)
        if t:
            out.append({"tipo": "parrafo", "texto": t, "fondo": fondo, "negrita": negrita})

    def tiene_bloque(el) -> bool:
        return any(isinstance(c.tag, str) and (c.tag.lower() in _BLOQUES or tiene_bloque(c)) for c in el)

    def recorrer(el, fondo_heredado=None):
        if not isinstance(el.tag, str):
            return
        tag = etree.QName(el).localname.lower() if "}" in el.tag else el.tag.lower()
        if tag in _OMITIR or "no-print" in _clases(el):
            return
        if tag == "svg":
            out.append({"tipo": "imagen", "svg": etree.tostring(el, encoding="unicode")})
            return
        if tag in ("h1", "h2", "h3", "h4"):
            nivel = {"h1": 0, "h2": 1, "h3": 2, "h4": 3}[tag]
            out.append({"tipo": "titulo", "nivel": nivel, "texto": _limpio(el.text_content()),
                        "fondo": fondo_heredado})
            return
        if tag == "table":
            filas = []
            for tr in el.iter("tr"):
                fondo_tr = _hex(_estilo(tr, "background"))
                celdas = []
                for td in tr:
                    if not isinstance(td.tag, str) or td.tag.lower() not in ("td", "th"):
                        continue
                    peso = (_estilo(td, "font-weight") or "") in ("bold", "700") or td.tag.lower() == "th"
                    celdas.append({"texto": _limpio(td.text_content()), "negrita": peso,
                                   "fondo": _hex(_estilo(td, "background")) or fondo_tr})
                if celdas:
                    filas.append(celdas)
            if filas:
                out.append({"tipo": "tabla", "filas": filas})
            return
        if tag in ("ul", "ol"):
            items = [_limpio(li.text_content()) for li in el.iter("li")]
            items = [i for i in items if i]
            if items:
                out.append({"tipo": "lista", "items": items})
            return
        fondo = _hex(_estilo(el, "background")) or fondo_heredado
        if tag in ("div", "p", "section", "body", "html", "pre") and tiene_bloque(el):
            parrafo(el.text, fondo)
            for c in el:
                recorrer(c, fondo)
                parrafo(c.tail, fondo)
            return
        negrita = (_estilo(el, "font-weight") or "") in ("bold", "700") or tag in ("strong", "b")
        parrafo(el.text_content(), None if fondo == "FFFFFF" else fondo, negrita)

    recorrer(raiz)
    # Las imágenes se convierten una sola vez y se reutilizan en Word y PDF.
    for b in out:
        if b["tipo"] == "imagen":
            try:
                b["png"] = svg_a_png(b["svg"], 1600)
            except Exception:
                b["png"] = None
    return out


# ── Word ─────────────────────────────────────────────────────────────────────
def _sombrear(elemento, hexa: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    props = elemento.get_or_add_tcPr() if hasattr(elemento, "get_or_add_tcPr") else elemento.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexa)
    props.append(shd)


def docx_desde_html(html: str, bloques: list[dict] | None = None) -> bytes:
    """Reporte como documento de Word editable (.docx), con las gráficas como imágenes."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Cm, Pt, RGBColor

    bloques = bloques if bloques is not None else bloques_desde_html(html)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)          # carta
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(1.8))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    ancho_util = sec.page_width - sec.left_margin - sec.right_margin

    for b in bloques:
        if b["tipo"] == "titulo":
            h = doc.add_heading(b["texto"], level=min(b["nivel"], 3) if b["nivel"] else 0)
            if b["nivel"] == 1:
                for r in h.runs:
                    r.font.color.rgb = RGBColor(0x1A, 0x56, 0x9A)
        elif b["tipo"] == "parrafo":
            p = doc.add_paragraph()
            run = p.add_run(b["texto"])
            run.bold = bool(b.get("negrita"))
            if b.get("fondo") and b["fondo"] != "FFFFFF" and not _es_oscuro(b["fondo"]):
                _sombrear(p._p, b["fondo"])
                run.font.size = Pt(9)
        elif b["tipo"] == "lista":
            for item in b["items"]:
                doc.add_paragraph(item, style="List Bullet")
        elif b["tipo"] == "imagen" and b.get("png"):
            doc.add_picture(io.BytesIO(b["png"]), width=ancho_util)
        elif b["tipo"] == "tabla":
            n_col = max(len(f) for f in b["filas"])
            t = doc.add_table(rows=0, cols=n_col)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for fila in b["filas"]:
                celdas = t.add_row().cells
                for j, c in enumerate(fila):
                    celdas[j].text = ""
                    run = celdas[j].paragraphs[0].add_run(c["texto"])
                    run.font.size = Pt(9)
                    run.bold = c["negrita"]
                    if c["fondo"] and c["fondo"] != "FFFFFF":
                        _sombrear(celdas[j]._tc, c["fondo"])
                        if _es_oscuro(c["fondo"]):
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                            run.bold = True
            doc.add_paragraph()
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────
def _fuentes_pdf() -> dict[str, str]:
    import matplotlib
    carpeta = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
    return {"": os.path.join(carpeta, "DejaVuSans.ttf"), "B": os.path.join(carpeta, "DejaVuSans-Bold.ttf"),
            "I": os.path.join(carpeta, "DejaVuSans-Oblique.ttf")}


_GLIFOS: set[int] | None = None


def _texto_pdf(texto: str) -> tuple[str, tuple[int, int, int] | None]:
    """Texto que la fuente del PDF puede dibujar y color del estado (🟢/🟠/🔴) si lo trae."""
    global _GLIFOS
    if _GLIFOS is None:
        from fontTools.ttLib import TTFont
        _GLIFOS = set(TTFont(_fuentes_pdf()[""])["cmap"].getBestCmap())
    color = None
    for emoji, rgb in _ESTADOS.items():
        if emoji in texto:
            color = color or rgb
            texto = texto.replace(emoji, "●")
    texto = "".join(ch for ch in texto if ord(ch) in _GLIFOS or ch in "\n")
    return re.sub(r"\s+", " ", texto).strip(), color


def pdf_desde_html(html: str, bloques: list[dict] | None = None) -> bytes:
    """Reporte en PDF (carta), con las mismas tablas, notas e imágenes."""
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    bloques = bloques if bloques is not None else bloques_desde_html(html)
    pdf = FPDF(format="letter")
    pdf.set_margins(15, 15, 15)
    pdf.set_auto_page_break(True, margin=15)
    for estilo, ruta in _fuentes_pdf().items():
        pdf.add_font("DejaVu", estilo, ruta)
    pdf.add_page()
    epw = pdf.epw

    for b in bloques:
        if b["tipo"] == "titulo":
            texto, _ = _texto_pdf(b["texto"])
            if not texto:
                continue
            if b["nivel"] == 1:
                pdf.ln(3)
                if pdf.get_y() > pdf.h - 45:
                    pdf.add_page()
                fondo = b.get("fondo") or "1A569A"
                pdf.set_fill_color(*(int(fondo[i:i + 2], 16) for i in (0, 2, 4)))
                pdf.set_text_color(255, 255, 255)
                pdf.set_font("DejaVu", "B", 11)
                pdf.multi_cell(epw, 7, texto, fill=True, new_x="LMARGIN", new_y="NEXT")
            else:
                pdf.set_text_color(26, 86, 154)
                pdf.set_font("DejaVu", "B", {0: 16, 2: 11, 3: 10}.get(b["nivel"], 10))
                pdf.multi_cell(epw, 6.5, texto, new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(44, 62, 80)
            pdf.ln(1)
        elif b["tipo"] == "parrafo":
            texto, color = _texto_pdf(b["texto"])
            if not texto:
                continue
            fondo = b.get("fondo")
            nota = bool(fondo and fondo != "FFFFFF" and not _es_oscuro(fondo))
            pdf.set_font("DejaVu", "B" if b.get("negrita") else "", 8.5 if nota else 9.5)
            pdf.set_text_color(*(color or (44, 62, 80)))
            if nota:
                pdf.set_fill_color(*(int(fondo[i:i + 2], 16) for i in (0, 2, 4)))
            pdf.multi_cell(epw, 4.6, texto, fill=nota, new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(44, 62, 80)
            pdf.ln(1.2)
        elif b["tipo"] == "lista":
            pdf.set_font("DejaVu", "", 9)
            for item in b["items"]:
                texto, color = _texto_pdf(item)
                pdf.set_text_color(*(color or (44, 62, 80)))
                pdf.multi_cell(epw, 4.6, "• " + texto, new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(44, 62, 80)
            pdf.ln(1)
        elif b["tipo"] == "imagen" and b.get("png"):
            from PIL import Image
            ancho_px, alto_px = Image.open(io.BytesIO(b["png"])).size
            alto = epw * alto_px / ancho_px
            if alto > pdf.h - 30:
                ancho = epw * (pdf.h - 30) / alto
                alto = pdf.h - 30
            else:
                ancho = epw
            if pdf.get_y() + alto > pdf.h - 15:
                pdf.add_page()
            pdf.image(io.BytesIO(b["png"]), x=pdf.l_margin + (epw - ancho) / 2, w=ancho, h=alto)
            pdf.ln(2)
        elif b["tipo"] == "tabla":
            n_col = max(len(f) for f in b["filas"])
            pdf.set_font("DejaVu", "B", 8)
            # Cada columna al menos tan ancha como su palabra más larga (los
            # números no se parten); el resto se reparte según el texto.
            minimos, pesos = [], []
            for j in range(n_col):
                textos = [_texto_pdf(f[j]["texto"])[0] if j < len(f) else "" for f in b["filas"]]
                palabras = [w for t in textos for w in t.split()] or [""]
                minimos.append(min(max(pdf.get_string_width(w) for w in palabras) + 3.0, epw * 0.4))
                pesos.append(min(max(max(len(t) for t in textos), 4), 45))
            libre = max(epw - sum(minimos), 0.0)
            anchos = [m + libre * p / sum(pesos) for m, p in zip(minimos, pesos)]
            pdf.set_font("DejaVu", "", 8)
            with pdf.table(col_widths=anchos, first_row_as_headings=False, line_height=4.2,
                           text_align="LEFT", borders_layout="HORIZONTAL_LINES") as tabla:
                for fila in b["filas"]:
                    r = tabla.row()
                    for j in range(n_col):
                        c = fila[j] if j < len(fila) else {"texto": "", "negrita": False, "fondo": None}
                        texto, color = _texto_pdf(c["texto"])
                        fondo = c["fondo"] if c["fondo"] and c["fondo"] != "FFFFFF" else None
                        oscuro = _es_oscuro(fondo)
                        estilo = FontFace(
                            emphasis="BOLD" if (c["negrita"] or oscuro) else None,
                            color=(255, 255, 255) if oscuro else (color or (44, 62, 80)),
                            # Siempre explícito: sin él fpdf2 rellena con el último color usado.
                            fill_color=tuple(int(fondo[i:i + 2], 16) for i in (0, 2, 4)) if fondo else (255, 255, 255),
                        )
                        r.cell(texto, style=estilo)
            pdf.ln(2)
    return bytes(pdf.output())


def documentos_reporte(html: str) -> dict[str, Any]:
    """Word y PDF del reporte, convirtiendo las gráficas una sola vez."""
    bloques = bloques_desde_html(html)
    return {"docx": docx_desde_html(html, bloques), "pdf": pdf_desde_html(html, bloques),
            "imagenes": sum(1 for b in bloques if b["tipo"] == "imagen" and b.get("png"))}
