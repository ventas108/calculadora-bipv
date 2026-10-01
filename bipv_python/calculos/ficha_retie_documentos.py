# -*- coding: utf-8 -*-
"""📋 Ficha RETIE en Word editable (.docx), PDF y PNG.

Spec ``07-informes/ficha-retie-word-pdf`` (1-oct-2026). La ficha solo se
descargaba en SVG: el PNG pedía CairoSVG (no está en el servidor) y Word no
muestra SVG. Igual que el 📄 Reporte (Spec ``07/reporte-word-pdf``), la ficha
se dibuja con Pillow (``calculos.svg_a_png``) y se arma:

- Word: la ficha a todo el ancho en una hoja carta horizontal y, en una hoja
  vertical, los datos del proyecto y la tabla de validaciones editable;
- PDF: lo mismo, con la fuente DejaVu.

Nada se recalcula: usa la misma configuración, cálculos y validaciones que
dibuja la ficha. Módulo puro: sin Streamlit.
"""
from __future__ import annotations

import io
from typing import Any

from calculos.reporte_documentos import _fuentes_pdf, _texto_pdf

_NIVEL = {
    "OK": ("OK", "EAF7F0", (22, 131, 93)),
    "PENDIENTE": ("Pendiente", "FFF3E3", (210, 107, 20)),
    "ERROR": ("Error", "FFF0F0", (197, 48, 48)),
}


def _datos_proyecto(cfg: dict, calc: dict, checks: list[dict]) -> list[tuple[str, str]]:
    p = cfg.get("proyecto", {})
    filas = [
        ("Proyecto", p.get("nombre_proyecto") or ""),
        ("Propietario", p.get("propietario") or ""),
        ("Dirección / municipio", ", ".join(x for x in (p.get("direccion"), p.get("municipio")) if x)),
        ("Operador de red", p.get("operador_red") or ""),
        ("Diseñador / matrícula", " · ".join(x for x in (p.get("disenador"), p.get("matricula")) if x)),
        ("Plano / revisión / fecha", " · ".join(str(x) for x in (p.get("plano"), p.get("revision"), p.get("fecha")) if x)),
    ]
    if calc.get("potencia_dc_kwp") is not None:
        filas.append(("Potencia DC", f"{calc['potencia_dc_kwp']:,.2f} kWp"))
    if calc.get("potencia_ac_kw") is not None:
        filas.append(("Potencia AC", f"{calc['potencia_ac_kw']:,.2f} kW"))
    n = {k: sum(1 for c in checks if c.get("nivel") == k) for k in _NIVEL}
    filas.append(("Validaciones", f"{n['OK']} OK · {n['PENDIENTE']} pendientes · {n['ERROR']} errores"))
    return [(a, b) for a, b in filas if b]


def docx_ficha_retie(png: bytes, cfg: dict, calc: dict, checks: list[dict]) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.shared import Cm, Pt, RGBColor

    from calculos.reporte_documentos import _sombrear

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(27.94), Cm(21.59)          # carta horizontal
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(1.0))
    nombre = (cfg.get("proyecto") or {}).get("nombre_proyecto") or "Proyecto"
    titulo = doc.add_paragraph()
    run = titulo.add_run(f"Ficha de validación RETIE — {nombre}")
    run.bold, run.font.size = True, Pt(12)
    run.font.color.rgb = RGBColor(0x1A, 0x56, 0x9A)
    doc.add_picture(io.BytesIO(png), width=sec.page_width - sec.left_margin - sec.right_margin
                    - Cm(0.1))

    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    sec2.orientation = WD_ORIENT.PORTRAIT
    sec2.page_width, sec2.page_height = Cm(21.59), Cm(27.94)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec2, lado, Cm(1.8))
    doc.add_heading("Datos del proyecto", level=1)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in _datos_proyecto(cfg, calc, checks):
        celdas = t.add_row().cells
        celdas[0].text, celdas[1].text = k, v
        celdas[0].paragraphs[0].runs[0].bold = True
    doc.add_heading("Validaciones", level=1)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for celda, texto in zip(t.rows[0].cells, ("Estado", "Validación", "Detalle")):
        celda.text = texto
        celda.paragraphs[0].runs[0].bold = True
        _sombrear(celda._tc, "1A569A")
        celda.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for c in checks:
        etiqueta, fondo, rgb = _NIVEL.get(c.get("nivel"), (c.get("nivel", ""), "FFFFFF", (44, 62, 80)))
        celdas = t.add_row().cells
        for celda, texto in zip(celdas, (etiqueta, c.get("titulo", ""), c.get("detalle", ""))):
            celda.text = ""
            r = celda.paragraphs[0].add_run(str(texto))
            r.font.size = Pt(9)
        celdas[0].paragraphs[0].runs[0].bold = True
        celdas[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(*rgb)
        _sombrear(celdas[0]._tc, fondo)
    for fila in t.rows:
        fila.cells[0].width, fila.cells[1].width, fila.cells[2].width = Cm(2.4), Cm(6.0), Cm(9.6)
    doc.add_paragraph(
        "Ficha preliminar para revisión técnica. Requiere memorias, coordinación de protecciones, "
        "estudio de cortocircuito, selección definitiva de equipos y firma de un profesional competente."
    ).runs[0].font.size = Pt(8.5)
    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


def pdf_ficha_retie(png: bytes, cfg: dict, calc: dict, checks: list[dict]) -> bytes:
    from fpdf import FPDF
    from fpdf.fonts import FontFace
    from PIL import Image

    pdf = FPDF(format="letter")
    for estilo, ruta in _fuentes_pdf().items():
        pdf.add_font("DejaVu", estilo, ruta)
    pdf.set_auto_page_break(True, margin=15)

    # Hoja 1: la ficha en horizontal, lo más grande que quepa.
    pdf.add_page(orientation="L")
    margen = 8.0
    ancho_px, alto_px = Image.open(io.BytesIO(png)).size
    ancho, alto = pdf.w - 2 * margen, (pdf.w - 2 * margen) * alto_px / ancho_px
    if alto > pdf.h - 2 * margen:
        alto = pdf.h - 2 * margen
        ancho = alto * ancho_px / alto_px
    pdf.image(io.BytesIO(png), x=(pdf.w - ancho) / 2, y=(pdf.h - alto) / 2, w=ancho, h=alto)

    # Hoja 2: datos y validaciones en texto.
    pdf.add_page(orientation="P")
    pdf.set_margins(15, 15, 15)
    pdf.set_xy(15, 15)
    epw = pdf.w - 30
    nombre, _ = _texto_pdf((cfg.get("proyecto") or {}).get("nombre_proyecto") or "Proyecto")
    pdf.set_font("DejaVu", "B", 14)
    pdf.set_text_color(26, 86, 154)
    pdf.multi_cell(epw, 7, f"Ficha de validación RETIE — {nombre}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    def _encabezado(texto):
        pdf.set_fill_color(26, 86, 154)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("DejaVu", "B", 11)
        pdf.multi_cell(epw, 7, texto, fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
        pdf.set_text_color(44, 62, 80)

    _encabezado("Datos del proyecto")
    pdf.set_font("DejaVu", "", 9)
    with pdf.table(col_widths=(epw * 0.32, epw * 0.68), first_row_as_headings=False, line_height=5,
                   text_align="LEFT", borders_layout="HORIZONTAL_LINES") as tabla:
        for k, v in _datos_proyecto(cfg, calc, checks):
            r = tabla.row()
            r.cell(_texto_pdf(k)[0], style=FontFace(emphasis="BOLD", fill_color=(255, 255, 255)))
            r.cell(_texto_pdf(v)[0], style=FontFace(fill_color=(255, 255, 255)))
    pdf.ln(3)

    _encabezado("Validaciones")
    pdf.set_font("DejaVu", "", 8.5)
    with pdf.table(col_widths=(epw * 0.13, epw * 0.32, epw * 0.55), line_height=4.4,
                   text_align="LEFT", borders_layout="HORIZONTAL_LINES",
                   headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255),
                                           fill_color=(26, 86, 154))) as tabla:
        cab = tabla.row()
        for texto in ("Estado", "Validación", "Detalle"):
            cab.cell(texto)
        for c in checks:
            etiqueta, fondo, rgb = _NIVEL.get(c.get("nivel"), (str(c.get("nivel", "")), "FFFFFF", (44, 62, 80)))
            r = tabla.row()
            r.cell(etiqueta, style=FontFace(emphasis="BOLD", color=rgb,
                                            fill_color=tuple(int(fondo[i:i + 2], 16) for i in (0, 2, 4))))
            r.cell(_texto_pdf(c.get("titulo", ""))[0], style=FontFace(fill_color=(255, 255, 255)))
            r.cell(_texto_pdf(c.get("detalle", ""))[0], style=FontFace(fill_color=(255, 255, 255)))
    pdf.ln(3)
    pdf.set_font("DejaVu", "I", 8)
    pdf.multi_cell(epw, 4, "Ficha preliminar para revisión técnica. Requiere memorias, coordinación de "
                   "protecciones, estudio de cortocircuito, selección definitiva de equipos y firma de un "
                   "profesional competente.", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def documentos_ficha_retie(svg: str, cfg: dict, calc: dict, checks: list[dict],
                           ancho_px: int = 2400) -> dict[str, Any]:
    """``{"png", "docx", "pdf"}`` de la ficha (bytes)."""
    from calculos.ficha_validacion_retie import exportar_ficha_png_bytes

    png = exportar_ficha_png_bytes(svg, ancho_px)
    return {"png": png, "docx": docx_ficha_retie(png, cfg, calc, checks),
            "pdf": pdf_ficha_retie(png, cfg, calc, checks)}
