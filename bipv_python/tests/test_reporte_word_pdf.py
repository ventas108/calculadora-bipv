# -*- coding: utf-8 -*-
"""Spec ``07-informes/reporte-word-pdf`` (1-oct-2026).

El usuario guardó el reporte de la Granja Apartadó en Word y se perdieron las
6 gráficas: Word no lee los SVG del HTML (el .docx quedó con 13 tablas y 0
imágenes). Ahora el reporte se descarga como Word editable (.docx) y PDF con
las gráficas como imágenes, armados del mismo HTML.
"""
import io
import re
import sys
from pathlib import Path

import pdfplumber
from docx import Document
from PIL import Image

from calculos.reporte_documentos import bloques_desde_html, docx_desde_html, documentos_reporte, pdf_desde_html
from calculos.svg_a_png import svg_a_png

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "tests"))
import test_reporte_granja_completo as granja  # noqa: E402

SVG = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 300">'
       '<rect x="0" y="0" width="200" height="300" fill="#FFFFFF"/>'
       '<rect x="10" y="250" width="40" height="40" fill="#3465A4"/>'
       '<polygon points="100,20 140,80 60,80" fill="rgb(230,145,56)"/>'
       '<polyline points="10,200 190,200" fill="none" stroke="#2E7D32" stroke-width="4" stroke-dasharray="6,4"/>'
       '<circle cx="150" cy="260" r="15" fill="#C0392B"/>'
       '<text x="100" y="150" font-size="20" text-anchor="middle" fill="#000">Fila 13 ★</text></svg>')


def _html():
    return granja._generar(**granja._estado())


# ── Gráficas a imagen ────────────────────────────────────────────────────────
def test_svg_a_png_dibuja_las_figuras_del_reporte():
    im = Image.open(io.BytesIO(svg_a_png(SVG, 400))).convert("RGB")
    assert im.size == (400, 600)                                   # respeta el viewBox (alto 300 → 600)
    assert im.getpixel((60, 540)) == (52, 101, 164)               # rect azul
    assert im.getpixel((200, 120)) == (230, 145, 56)              # triángulo naranja
    assert im.getpixel((300, 520)) == (192, 57, 43)               # círculo rojo
    assert len(set(im.crop((100, 280, 300, 320)).getdata())) > 2  # texto dibujado


def test_svg_con_viewbox_en_minusculas():
    # El parser de HTML pasa los atributos a minúsculas; las figuras altas no se cortan.
    im = Image.open(io.BytesIO(svg_a_png(SVG.replace("viewBox", "viewbox"), 400)))
    assert im.size == (400, 600)


# ── Bloques, Word y PDF ──────────────────────────────────────────────────────
def test_bloques_del_reporte_de_granja():
    html = _html()
    bloques = bloques_desde_html(html)
    n_svg = len(re.findall(r"<svg", html))
    imagenes = [b for b in bloques if b["tipo"] == "imagen"]
    assert n_svg >= 5 and len(imagenes) == n_svg and all(b["png"] for b in imagenes)
    titulos = [b["texto"] for b in bloques if b["tipo"] == "titulo"]
    assert any("Información General del Proyecto" in t for t in titulos)
    assert any("Granja FV" in t for t in titulos)
    assert sum(b["tipo"] == "tabla" for b in bloques) >= 8


def test_word_editable_con_todas_las_graficas():
    html = _html()
    doc = Document(io.BytesIO(docx_desde_html(html)))
    assert len(doc.inline_shapes) == len(re.findall(r"<svg", html))
    assert len(doc.tables) >= 8
    texto = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)
    for t in ("Información General del Proyecto", "Área del terreno", "Calidad del módulo (aplicada)",
              "Plano eléctrico visto desde arriba", "INV-1", "Paso de la maquinaria agrícola"):
        assert t in texto, t
    # carta con márgenes de 1,8 cm
    sec = doc.sections[0]
    assert round(sec.page_width.cm, 1) == 21.6 and round(sec.left_margin.cm, 1) == 1.8


def test_pdf_con_texto_tablas_e_imagenes():
    html = _html()
    with pdfplumber.open(io.BytesIO(pdf_desde_html(html))) as pdf:
        texto = "\n".join(p.extract_text() or "" for p in pdf.pages)
        imagenes = sum(len(p.images) for p in pdf.pages)
    assert imagenes == len(re.findall(r"<svg", html))
    for t in ("Información General del Proyecto", "Área del terreno", "Calidad del módulo", "INV-1"):
        assert t in texto, t
    assert "PVsyst" not in texto


def test_documentos_juntos():
    d = documentos_reporte(_html())
    assert d["docx"][:2] == b"PK" and d["pdf"][:4] == b"%PDF" and d["imagenes"] >= 5


# ── Página ───────────────────────────────────────────────────────────────────
def test_la_pagina_ofrece_word_y_pdf():
    import calculos.auth as _auth
    import calculos.trm_utils as _trm
    from streamlit.testing.v1 import AppTest
    _auth.requerir_login = lambda solo_admin=False: {"email": "t@t", "rol": "admin", "activo": True, "nombre": "T"}
    at = AppTest.from_file(str(granja.PAGINA), default_timeout=180)
    base = {"recurso_solar_ok": True, "produccion_ok": True, "res_produccion": granja.RES,
            "poa_anual_kWh_m2": 1838.0, "nombre_proyecto": "Granja Apartadó",
            _trm._KEY_VALOR: 3900.0, _trm._KEY_FUENTE: "manual"}
    for k, v in {**base, **granja._estado()}.items():
        at.session_state[k] = v
    at.run()
    next(b for b in at.button if "Generar Reporte" in b.label).click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.session_state["_reporte_docx"][:2] == b"PK"
    assert at.session_state["_reporte_pdf"][:4] == b"%PDF"
    src = granja.PAGINA.read_text(encoding="utf-8")
    assert "Word editable (.docx)" in src and "print-color-adjust:exact" in src


def test_manual_del_asistente_lo_explica():
    kb = (RAIZ / "datos" / "base_conocimiento_asistente.md").read_text(encoding="utf-8")
    seccion = kb[kb.index("## 106."):]
    for texto in ("Word editable", "PDF", "gráficas", "No abras el HTML con Word", "Gráficos de fondo"):
        assert texto in seccion, texto
