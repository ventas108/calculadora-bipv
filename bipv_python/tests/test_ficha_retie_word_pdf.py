# -*- coding: utf-8 -*-
"""Spec 07-informes/ficha-retie-word-pdf (1-oct-2026).

La 📋 Ficha RETIE solo se descargaba en SVG: el PNG pedía CairoSVG, que no
está en el servidor, y Word no muestra SVG. Ahora sale en PNG, Word editable
y PDF con el mismo dibujo (Pillow) y las validaciones como texto.
"""
import io
from pathlib import Path

from PIL import Image

from calculos.ficha_validacion_retie import (
    calcular_retie,
    construir_config_retie,
    exportar_ficha_png_bytes,
    generar_ficha_svg,
    validar_retie,
)
from calculos.ficha_retie_documentos import documentos_ficha_retie
from calculos.svg_a_png import svg_a_png

_RAIZ = Path(__file__).resolve().parents[1]


def _ficha():
    cfg = construir_config_retie(
        nombre_proyecto="Granja Apartadó", panel_nombre="JAM66D46-720/LB", potencia_w=720.0,
        voc_v=49.0, vmp_v=41.0, isc_a=18.59, coef_voc_pct_c=-0.25, inversor_nombre="Growatt MAX 100KTL3 LV",
        potencia_ac_kw_unidad=100.0, n_inversores=2, tension_salida_v=400.0, vdc_max_v=1100.0,
        vmppt_min_v=180.0, vmppt_max_v=1000.0, n_paneles=300, n_serie=20, strings_por_inversor=[8, 7],
        temperatura_minima_diseno_c=18.0,
    )
    calc = calcular_retie(cfg)
    checks = validar_retie(cfg, calc)
    return cfg, calc, checks, generar_ficha_svg(cfg, calc, checks)


def _pixeles_de_color(png: bytes) -> int:
    img = Image.open(io.BytesIO(png)).convert("RGB").resize((450, 355))
    return sum(1 for r, g, b in img.getdata() if max(r, g, b) - min(r, g, b) > 60)


def test_png_sin_cairosvg_con_el_tamano_y_los_colores_de_la_ficha():
    _, _, _, svg = _ficha()
    png = exportar_ficha_png_bytes(svg)          # en este entorno no hay CairoSVG
    assert png and png[:8] == b"\x89PNG\r\n\x1a\n"
    img = Image.open(io.BytesIO(png))
    assert img.size[0] == 2400 and abs(img.size[1] / img.size[0] - 1420 / 1800) < 0.01
    assert _pixeles_de_color(png) > 2000          # tarjetas y bloques de color, no solo texto


def test_path_y_defs():
    svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">'
           '<defs><marker id="a"><path d="M0,0 L0,60 L90,30 z" fill="#ff0000"/></marker></defs>'
           '<path d="M10 90 V50 H90 V90" fill="none" stroke="#0000ff" stroke-width="4"/></svg>')
    img = Image.open(io.BytesIO(svg_a_png(svg, 100))).convert("RGB")
    assert img.getpixel((10, 70))[2] > 200 and img.getpixel((50, 50))[2] > 200   # el path se dibuja
    assert img.getpixel((5, 20)) == (255, 255, 255)                               # lo de <defs> no


def test_word_y_pdf_con_la_ficha_y_las_validaciones_en_texto():
    cfg, calc, checks, svg = _ficha()
    docs = documentos_ficha_retie(svg, cfg, calc, checks)
    from docx import Document
    doc = Document(io.BytesIO(docs["docx"]))
    assert len(doc.inline_shapes) == 1
    texto = " ".join(c.text for t in doc.tables for f in t.rows for c in f.cells)
    for c in checks:
        assert c["titulo"] in texto
    assert "Granja Apartadó" in " ".join(p.text for p in doc.paragraphs)
    # Primera sección horizontal para que la ficha se lea a buen tamaño.
    assert doc.sections[0].page_width > doc.sections[0].page_height

    assert docs["pdf"][:5] == b"%PDF-"
    _assert_pdf(docs["pdf"], checks)


def _assert_pdf(pdf: bytes, checks):
    import pdfplumber
    with pdfplumber.open(io.BytesIO(pdf)) as d:
        assert len(d.pages) >= 2
        assert d.pages[0].width > d.pages[0].height             # ficha en horizontal
        assert len(d.pages[0].images) == 1
        texto = " ".join(p.extract_text() or "" for p in d.pages)
    assert "Validaciones" in texto and checks[0]["titulo"] in texto


def test_la_pagina_ofrece_word_pdf_png_y_svg():
    src = (_RAIZ / "pages" / "21_📋_Ficha_Validacion_RETIE.py").read_text(encoding="utf-8")
    assert "documentos_ficha_retie(" in src
    for etiqueta in ("⬇️ Word editable", "⬇️ PDF", "⬇️ PNG", "⬇️ SVG"):
        assert etiqueta in src
    assert "requiere CairoSVG" not in src


def test_manual_del_asistente_lo_explica():
    kb = (_RAIZ / "datos" / "base_conocimiento_asistente.md").read_text(encoding="utf-8")
    i = kb.index("## 110.")
    s = kb[i:kb.find("\n## ", i + 5) if kb.find("\n## ", i + 5) > 0 else None]
    for t in ("📋 Ficha RETIE", "Word editable", "PDF", "PNG", "SVG"):
        assert t in s
    assert i < kb.rindex("Calculadora BIPV — Innovación Química")
