"""Genera el Anexo 72 de auditoría en Word desde su fuente Markdown.

Ejecutar desde cualquier directorio del repositorio:

    bipv_python/.venv/bin/python scripts/generar_anexo_72_auditoria_docx.py
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SOURCE_PATH = ROOT / "bipv_python" / "datos" / "base_conocimiento_asistente.md"
OUT_PATH = ROOT / "Anexo_72_Esquema_Practico_Auditoria_BIPV.docx"
SECTION_PATTERN = re.compile(
    r"^## 72\.\s+(?P<title>.+?)\n\n(?P<body>.*?)(?=^##\s|\Z)",
    re.MULTILINE | re.DOTALL,
)
INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def _add_inline(paragraph, text: str) -> None:
    """Convierte negritas y código Markdown simple en runs de Word."""
    for part in INLINE_PATTERN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def _add_body(doc: Document, body: str) -> None:
    for block in re.split(r"\n\s*\n", body.strip()):
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        for line in lines:
            numbered = re.match(r"^\d+\.\s+(.+)$", line)
            bullet = re.match(r"^-\s+(.+)$", line)
            if numbered:
                paragraph = doc.add_paragraph(style="List Number")
                _add_inline(paragraph, numbered.group(1))
            elif bullet:
                paragraph = doc.add_paragraph(style="List Bullet")
                _add_inline(paragraph, bullet.group(1))
            else:
                paragraph = doc.add_paragraph()
                _add_inline(paragraph, line)


def construir() -> Path:
    source = SOURCE_PATH.read_text(encoding="utf-8")
    match = SECTION_PATTERN.search(source)
    if not match:
        raise RuntimeError(f"No se encontró el Anexo 72 en {SOURCE_PATH}")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    source_title = match.group("title").strip()
    display_title = source_title.replace("Anexo —", "Anexo 72 —", 1)
    title = doc.add_heading(display_title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Guía operativa sincronizada con el Director SDD y la base de conocimiento del asistente"
    )
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    _add_body(doc, match.group("body"))

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(
        "Calculadora BIPV — Innovación Química · Actualizado 19-sep-2026"
    )
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.core_properties.title = "Anexo 72 — Esquema práctico de auditoría BIPV"
    doc.core_properties.subject = "Auditoría operativa SDD de proyectos BIPV"
    doc.save(OUT_PATH)
    print(f"Generado: {OUT_PATH}")
    return OUT_PATH


if __name__ == "__main__":
    construir()
