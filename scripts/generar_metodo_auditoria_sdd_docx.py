"""Genera Metodo_Auditoria_SDD_Calculadora_BIPV.docx a partir del método de
auditoría SDD actualizado el 19-sep-2026. Ejecutar con el mismo entorno que ya
tiene python-docx (requirements.txt del proyecto):

    cd bipv_python && .venv/bin/python ../scripts/generar_metodo_auditoria_sdd_docx.py

El archivo se crea en la raíz del repositorio.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_PATH = Path(__file__).resolve().parent.parent / "Metodo_Auditoria_SDD_Calculadora_BIPV.docx"


def _heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def _p(doc, text, bold=False, italic=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    return par


def _table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    return t


def construir():
    doc = Document()

    title = doc.add_heading(
        "Método de auditoría SDD — Calculadora BIPV", level=0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _p(
        doc,
        "Referencia para verificar cada corrida de un proyecto real contra "
        "lo implementado en CodeSpecs. Complementa el documento de memoria "
        "'Flujo operativo de confianza' (19-sep-2026).",
        italic=True,
    )

    _heading(doc, "Paso 0 — Punto de entrada: el director", level=1)
    _p(doc, "Revisar en este orden antes de tocar la app:")
    for item in [
        "CodeSpecs/00-director/vision.md — principio: nada se marca "
        "completado sin validación real.",
        "CodeSpecs/00-director/arquitectura-global.md — propiedad de cada "
        "motor y funciones exclusivas de React o Streamlit.",
        "CodeSpecs/00-director/mapa-dependencias.md — orden declarado "
        "(01→09) y excepciones de runtime ya documentadas (ej. 05→04).",
        "CodeSpecs/00-director/contratos-entre-modulos.md — contrato "
        "vigente por módulo, comparadores y asistentes; incluye desviaciones activas.",
        "CodeSpecs/00-director/registro-de-decisiones.md — historial "
        "cronológico de por qué se corrigió cada cosa; sirve para detectar "
        "si una garantía fue des-hecha después.",
        "CodeSpecs/00-director/separacion-apps.md — frontera obligatoria entre "
        "React y Streamlit, rutas, procesos y despliegues independientes.",
    ]:
        doc.add_paragraph(item, style="List Number")

    _heading(doc, "Paso 1 — Por cada módulo (01-09), leer solo 3 cosas", level=1)
    _table(
        doc,
        ["Archivo", "Qué te dice"],
        [
            ["problema.md (línea Estado)",
             "completado / en validación / idea"],
            ["diseno.md → \"Criterios de aceptación\"",
             "El checklist de auditor: lo que el módulo promete"],
            ["validacion.md → checklist + Resultado",
             "Evidencia real que respalda la promesa y qué falta"],
        ],
    )
    _p(
        doc,
        "Estado real al 19-sep-2026: 01 y 03–08 completados; 02 raíz archivado "
        "y reemplazado por Specs verticales (React aprobado, Streamlit en idea); "
        "09-despliegue en validación. La independencia de Streamlit no la deja "
        "fuera del Director: exige contrato y despliegue propios.",
    )

    _heading(
        doc, "Paso 2 — Traducir cada criterio a una señal visible en la app",
        level=1,
    )
    _table(
        doc,
        ["Módulo", "Criterio de aceptación (resumen)", "Señal en la app"],
        [
            ["03", "Diseño no vigente no produce resultados persistibles",
             "Banner vigente/aviso en Dimensionamiento"],
            ["04", "Simulación exige compatibilidad Y vigencia",
             "Bloqueo explícito del botón Simular"],
            ["05", "Un solo cálculo térmico, sin doble conteo",
             "Banner \"SDM usa POA sin térmico + k_BIPV\""],
            ["06", "Restauración solo si el payload verifica",
             "Aviso \"📂 Datos restaurados...\" o su ausencia"],
            ["07", "Informes hereda vigencia, no la verifica",
             "Sin alarma propia — confiar solo si 04/06 ya validaron"],
            ["08", "La infraestructura UI no invalida; adoptar sí es una operación explícita",
             "Solo los botones Adoptar aplican la invalidación central definida"],
            ["09", "Servidor en el mismo commit que GitHub",
             "git rev-parse --short HEAD idéntico en ambos lados"],
        ],
    )

    _heading(
        doc, "Paso 3 — Auditar los comparadores Streamlit", level=1
    )
    _table(
        doc,
        ["Comparador", "Invariante", "Comprobación"],
        [
            ["Inversores",
             "Cada candidato aplica su clipping a la potencia previa al recorte",
             "Si falta P_ac_sin_recorte_kW, debe exigir volver a Producción"],
            ["Paneles",
             "✅ adoptable; ❌ incompatible; — no evaluable. Adopta la ficha simulada",
             "Un resultado legacy se descarta y se repite; nunca se reconstruye"],
            ["Orientación",
             "No cambia compatibilidad eléctrica y recalcula POA antes de adoptar",
             "Después de adoptar deben caducar Producción, Finanzas y CO₂"],
            ["Asistentes",
             "El general usa manual+estado; cada Analista local usa solo su tabla actual",
             "Ningún agente modifica el proyecto ni adopta alternativas"],
        ],
    )

    _heading(doc, "Paso 4 — Revisar desviaciones activas", level=1)
    _p(
        doc,
        "No certificar estos puntos como resueltos hasta cerrar una Spec vertical:",
    )
    for item in [
        "Orientación multi-superficie: la adopción global puede eliminar estado "
        "multi-superficie; falta decidir bloqueo o adopción por superficie.",
        "Vigencia de tablas e IA: los resultados de los comparadores todavía no "
        "comparten una firma de entradas que pruebe que siguen actuales.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    _heading(
        doc, "Paso 5 — Ritual antes de correr un proyecto de cliente", level=1
    )
    for item in [
        "Verificar 09: commit idéntico en GitHub y servidor.",
        "Confirmar en registro-de-decisiones.md que nada reabre/revierte lo "
        "que se va a asumir como cerrado.",
        "Correr el proyecto completo en orden, en una sesión limpia.",
        "Verificar que cada alarma prometida aparece cuando corresponde.",
        "Si algo no coincide con su CodeSpec: es una incoherencia real, se "
        "documenta (problema.md nuevo o reabierto), no se ignora.",
    ]:
        doc.add_paragraph(item, style="List Number")

    _heading(doc, "Límites honestos del método", level=1)
    for item in [
        "No valida que el resultado financiero específico de un cliente sea "
        "correcto — solo garantiza que los datos usados son vigentes.",
        "07 no tiene alarma propia por diseño. La infraestructura común de 08 "
        "tampoco invalida; los comandos explícitos de adopción son la excepción.",
        "El agente SDD recibe la Spec y los seis documentos del Director, pero no "
        "recibe automáticamente rama, diff ni estado Git: se verifican aparte.",
        "El agente SDD prepara, detecta riesgos y bloquea; no aprueba, implementa "
        "ni sustituye las pruebas reales.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Calculadora BIPV — Innovación Química · Actualizado 19-sep-2026")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.save(OUT_PATH)
    print(f"Generado: {OUT_PATH}")


if __name__ == "__main__":
    construir()
